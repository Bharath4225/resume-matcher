import io
import os
import PyPDF2
import docx

def extract_text_from_file(file_storage) -> str:
    """
    Extracts plain text from uploaded FileStorage objects.
    Supports .txt, .pdf, and .docx formats.
    """
    filename = file_storage.filename.lower()
    content_bytes = file_storage.read()
    
    # Reset buffer position for safety
    file_storage.seek(0)
    
    if filename.endswith(".pdf"):
        return extract_text_from_pdf(content_bytes)
    elif filename.endswith(".docx"):
        return extract_text_from_docx(content_bytes)
    elif filename.endswith(".txt"):
        return extract_text_from_txt(content_bytes)
    else:
        # Fallback: attempt UTF-8 decoding
        return extract_text_from_txt(content_bytes)

def extract_text_from_txt(content_bytes: bytes) -> str:
    """Decodes plain text with fallbacks for encodings."""
    for encoding in ['utf-8', 'latin-1', 'cp1252']:
        try:
            return content_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content_bytes.decode('utf-8', errors='ignore')

def extract_text_from_pdf(content_bytes: bytes) -> str:
    """Extracts text from PDF binary using PyPDF2."""
    try:
        pdf_file = io.BytesIO(content_bytes)
        reader = PyPDF2.PdfReader(pdf_file)
        text_runs = []
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text_runs.append(extracted)
        return "\n".join(text_runs)
    except Exception as e:
        raise ValueError(f"Failed to parse PDF file: {str(e)}")

def extract_text_from_docx(content_bytes: bytes) -> str:
    """Extracts text from DOCX binary using python-docx."""
    try:
        docx_file = io.BytesIO(content_bytes)
        doc = docx.Document(docx_file)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        full_text.append(cell.text)
        return "\n".join(full_text)
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX file: {str(e)}")
